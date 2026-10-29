"""把 docs/thesis 下的章节 Markdown 组装为本科毕业论文 Word。

用法: .venv/bin/python scripts/build_thesis_docx.py
输出: docs/thesis/毕业论文-企业大模型数据防泄漏系统.docx

目录与页码是 Word 域；首次打开若未自动刷新，Ctrl+A 后 F9 更新。
"""
import os
import re

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D = os.path.join(ROOT, "docs", "thesis")
OUT = os.path.join(D, "毕业论文-企业大模型数据防泄漏系统.docx")
FIGS = os.path.join(D, "figures")

SERIF = "宋体"
SANS = "黑体"
LATIN = "Times New Roman"
MONO = "Consolas"
TEXT_WIDTH = Cm(15.5)
MAX_FIG_H = Cm(18.0)
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def set_run_font(run, size=12, bold=None, latin=LATIN, east=SERIF, color=None):
    run.font.name = latin
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)


def set_style_font(style, size, bold=False, latin=LATIN, east=SERIF):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)


def add_inline(p, text, size=12, bold=False, east=SERIF, latin=LATIN, color=None):
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            set_run_font(p.add_run(text[pos:m.start()]), size, bold, latin, east, color)
        token = m.group(0)
        if token.startswith("**"):
            set_run_font(p.add_run(token[2:-2]), size, True, latin, east, color)
        else:
            set_run_font(p.add_run(token[1:-1]), max(8, size - 1), False, MONO, east, color)
        pos = m.end()
    if pos < len(text):
        set_run_font(p.add_run(text[pos:]), size, bold, latin, east, color)


def body(text, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, east=SERIF, latin=LATIN):
    p = doc.add_paragraph(style="Normal")
    pf = p.paragraph_format
    pf.first_line_indent = Pt(24) if indent else Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = align
    add_inline(p, text, size, east=east, latin=latin)
    return p


def heading(text, level, page_break=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    if page_break:
        p.paragraph_format.page_break_before = True
    add_inline(p, text, {1: 16, 2: 14, 3: 12}[level], east=SANS)
    return p


def toc_title():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.page_break_before = True
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(18)
    set_run_font(p.add_run("目　　录"), 16, False, LATIN, SANS)


def caption(text, table=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.keep_with_next = table
    add_inline(p, text, 10.5, bold=True)


def note(text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.74)
    pf.first_line_indent = Pt(0)
    pf.line_spacing = 1.25
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    add_inline(p, text, 10.5, color=RGBColor(0x59, 0x59, 0x59))


def list_item(text, bullet):
    p = doc.add_paragraph(style="Normal")
    pf = p.paragraph_format
    pf.left_indent = Cm(0.74)
    pf.first_line_indent = Cm(-0.74)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(p, ("• " + text) if bullet else text)


def code_block(lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.first_line_indent = Pt(0)
    pf.line_spacing = 1.0
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        set_run_font(p.add_run(line.replace("\t", "    ")), 9, False, MONO)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F2F2F2")
    ppr = p._p.get_or_add_pPr()
    anchor = ppr.find(qn("w:spacing"))
    if anchor is None:
        anchor = ppr.find(qn("w:ind"))
    if anchor is not None:
        anchor.addprevious(shd)
    else:
        ppr.append(shd)


def image(src, cap):
    path = os.path.join(D, src)
    if not os.path.exists(path):
        print("缺图:", path)
        body(cap or src)
        return
    with Image.open(path) as im:
        w, h = im.size
    width, height = TEXT_WIDTH, Cm(15.5 * h / max(1, w))
    max_h = Cm(18.0)
    if height > max_h:
        height = max_h
        width = Cm(max_h.cm * w / max(1, h))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)
    pf.keep_with_next = True
    p.add_run().add_picture(path, width=width, height=height)
    if cap:
        caption(cap)
        caption_p = doc.paragraphs[-1]
        caption_p.paragraph_format.keep_together = True


def len_disp(text):
    return sum(2 if ord(ch) > 127 else 1 for ch in text)


def table(block):
    rows = []
    for line in block:
        s = line.strip()
        if re.fullmatch(r"\|[\s:\-|]+\|", s):
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        rows.append(cells)
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    for el in tblPr.findall(qn("w:tblW")):
        tblPr.remove(el)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), "8788")
    anchor = tblPr.find(qn("w:jc"))
    if anchor is None:
        anchor = tblPr.find(qn("w:tblLayout"))
    if anchor is None:
        anchor = tblPr.find(qn("w:tblLook"))
    if anchor is not None:
        anchor.addprevious(tblW)
    else:
        tblPr.append(tblW)
    weights = []
    for ci in range(ncol):
        longest = max(2, max(len_disp(r[ci]) for r in rows))
        weights.append(min(40, longest))
    base = 1.5
    extra = max(0.0, 15.5 - base * ncol)
    widths = [base + extra * w / sum(weights) for w in weights]
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 567)))
    for ri, row in enumerate(rows):
        tr = t.rows[ri]._tr
        trPr = tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            hdr = OxmlElement("w:tblHeader")
            hdr.set(qn("w:val"), "true")
            trPr.append(hdr)
        for ci, text in enumerate(row):
            cell = t.cell(ri, ci)
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.first_line_indent = Pt(0)
            pf.space_before = Pt(1)
            pf.space_after = Pt(1)
            pf.line_spacing = 1.0
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, text, 10.5, bold=ri == 0)
    doc.add_paragraph()


def smart_join(lines):
    text = " ".join(lines)
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", text)
    text = re.sub(r"\s+(?=[，。；：、！？）》”])", "", text)
    text = re.sub(r"(?<=[，。；：、！？）》”])\s+", "", text)
    text = re.sub(r"\s+(?=[（《“])", "", text)
    text = re.sub(r"(?<=[（《“])\s+", "", text)
    return text


def render(lines, ctx):
    i, para, code = 0, [], []
    in_code = False
    seen_h2 = False
    while i < len(lines):
        line = lines[i]
        if in_code:
            if line.strip().startswith("```"):
                code_block(code)
                code, in_code = [], False
            else:
                code.append(line)
            i += 1
            continue
        if line.strip().startswith("```"):
            if para:
                flush(para, ctx)
                para = []
            code, in_code = [], True
            i += 1
            continue
        if not line.strip():
            if para:
                flush(para, ctx)
                para = []
            i += 1
            continue
        if line.startswith("#"):
            if para:
                flush(para, ctx)
                para = []
            level = len(line) - len(line.lstrip("#"))
            if 1 <= level <= 3:
                text = line[level:].strip()
                if level == 1:
                    heading(text, 1, page_break=not ctx["first_h1"])
                    ctx["first_h1"] = False
                else:
                    heading(text, level)
                    if level == 2:
                        seen_h2 = True
            i += 1
            continue
        if line.lstrip().startswith(">"):
            if para:
                flush(para, ctx)
                para = []
            if ctx["skip_meta"] and not seen_h2:
                while i < len(lines) and lines[i].lstrip().startswith(">"):
                    i += 1
                ctx["skip_meta"] = False
                continue
            block = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                block.append(lines[i].lstrip()[1:].strip())
                i += 1
            note(smart_join(block))
            continue
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line.strip())
        if m:
            if para:
                flush(para, ctx)
                para = []
            alt, src = m.group(1), m.group(2)
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            cap = alt
            if j < len(lines) and re.fullmatch(r"\*\*图 .+\*\*", lines[j].strip()):
                cap = lines[j].strip().strip("*")
                i = j
            image(src, cap)
            i += 1
            continue
        if line.startswith("|"):
            if para:
                flush(para, ctx)
                para = []
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            table(block)
            continue
        if re.fullmatch(r"\*\*[表图] .+\*\*", line.strip()):
            if para:
                flush(para, ctx)
                para = []
            caption(line.strip().strip("*"), table=line.strip().startswith("**表"))
            i += 1
            continue
        if line.startswith("- ") or re.match(r"^\d+\. ", line):
            if para:
                flush(para, ctx)
                para = []
            bullet = line.startswith("- ")
            item = [line[2:].strip() if bullet else line.strip()]
            i += 1
            while i < len(lines) and lines[i].strip() and (
                    lines[i].startswith("  ") or lines[i].startswith("\t")):
                item.append(lines[i].strip())
                i += 1
            list_item(smart_join(item), bullet)
            continue
        if line.strip() == "---":
            if para:
                flush(para, ctx)
                para = []
            i += 1
            continue
        para.append(line.strip())
        i += 1
    if para:
        flush(para, ctx)
    if in_code and code:
        code_block(code)


def flush(lines, ctx):
    text = smart_join(lines)
    if not text:
        return
    if text.startswith("**关键词**") or text.startswith("**Keywords**"):
        body(text, indent=False)
    elif ctx.get("en"):
        body(text, indent=False, latin=LATIN)
    else:
        body(text)


def cover():
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    set_run_font(p.add_run("（学校名称）"), 18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(12)
    set_run_font(p.add_run("本科毕业设计（论文）"), 24, False, LATIN, SANS)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    set_run_font(p.add_run("基于规则与小模型级联的企业大模型\n数据防泄漏系统设计与实现"), 22, True, LATIN, SANS)
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_table(rows=6, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = ["学　　院", "专　　业", "学生姓名", "学　　号", "指导教师", "完成日期"]
    for i, label in enumerate(labels):
        c0, c1 = t.cell(i, 0), t.cell(i, 1)
        c0.width, c1.width = Cm(3.5), Cm(8.0)
        p0, p1 = c0.paragraphs[0], c1.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p0.add_run(label + "："), 14)
        set_run_font(p1.add_run("________________"), 14)
        for p in (p0, p1):
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    set_run_font(p.add_run("2026 年　　月"), 14)


def add_field(p, instr, placeholder, size=10.5):
    def field_run():
        run = p.add_run()
        set_run_font(run, size)
        return run
    r = field_run()
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), "begin")
    r._r.append(el)
    r = field_run()
    el = OxmlElement("w:instrText")
    el.set(qn("xml:space"), "preserve")
    el.text = instr
    r._r.append(el)
    r = field_run()
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), "separate")
    r._r.append(el)
    run = p.add_run(placeholder)
    set_run_font(run, size)
    r = field_run()
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), "end")
    r._r.append(el)


def add_page_field(section, fmt, start):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    add_field(p, " PAGE ", "1")
    sectPr = section._sectPr
    for el in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(el)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))
    cols = sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(pg)
    else:
        sectPr.append(pg)


def _strip_md(text):
    return re.sub(r"\*\*(.+?)\*\*", r"\1", re.sub(r"`([^`]+)`", r"\1", text))


def collect_toc_headings():
    """预扫描章节 Markdown 收集标题，作为 TOC 域的缓存结果（Word 更新域后自动替换）。"""
    heads = [(1, "摘要"), (1, "Abstract")]
    files = ("13-ch1-introduction.md", "14-ch2-literature-review.md",
             "11-ch3-requirements.md", "12-ch4-design.md",
             "08-ch5-implementation.md", "07-ch6-testing.md")

    def scan_lines(lines, out):
        in_code = False
        for line in lines:
            if line.strip().startswith("```"):
                in_code = not in_code
                continue
            if in_code:
                continue
            m = re.match(r"^(#{1,3})\s+(.*)", line)
            if m and 1 <= len(m.group(1)) <= 3:
                out.append((len(m.group(1)), _strip_md(m.group(2).strip())))

    for name in files:
        scan_lines(read(name), heads)
    s = open(os.path.join(D, "09-abstract-conclusion.md"), encoding="utf-8").read()
    scan_lines(s[s.index("# 第 7 章"):].splitlines(), heads)
    heads += [(1, "参考文献"), (1, "致谢")]
    scan_lines(read("16-appendix.md"), heads)
    return heads


def add_toc():
    heads = collect_toc_headings()
    for i, (lvl, text) in enumerate(heads):
        p = doc.add_paragraph(style=f"TOC {lvl}")
        pf = p.paragraph_format
        pf.first_line_indent = Pt(0)
        if i == 0:
            r = p.add_run()
            set_run_font(r, 12)
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), "begin")
            r._r.append(el)
            r = p.add_run()
            set_run_font(r, 12)
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = ' TOC \\o "1-3" \\h \\z \\u '
            r._r.append(el)
            r = p.add_run()
            set_run_font(r, 12)
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), "separate")
            r._r.append(el)
        size = 12 if lvl <= 2 else 10.5
        set_run_font(p.add_run(text), size)
        r = p.add_run("\t")
        set_run_font(r, size)
        if i == len(heads) - 1:
            r = p.add_run()
            set_run_font(r, 12)
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), "end")
            r._r.append(el)


def toc_styles():
    for name, size, indent in (("TOC 1", 12, 0), ("TOC 2", 12, 0.74), ("TOC 3", 10.5, 1.48)):
        try:
            st = doc.styles[name]
        except KeyError:
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(st, size, False, LATIN, SERIF)
        pf = st.paragraph_format
        pf.left_indent = Cm(indent)
        pf.first_line_indent = Pt(0)
        pf.line_spacing = 1.5 if name != "TOC 3" else 1.25
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def setup_doc():
    global doc
    doc = Document()
    doc.core_properties.title = "基于规则与小模型级联的企业大模型数据防泄漏系统设计与实现"
    doc.core_properties.subject = "本科毕业设计（论文）"
    doc.core_properties.author = ""
    normal = doc.styles["Normal"]
    set_style_font(normal, 12, False, LATIN, SERIF)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(24)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for i, (size, after) in enumerate(((16, 18), (14, 6), (12, 3)), 1):
        st = doc.styles[f"Heading {i}"]
        set_style_font(st, size, False, LATIN, SANS)
        st.paragraph_format.first_line_indent = Pt(0)
        st.paragraph_format.space_before = Pt(0 if i == 1 else 12)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 1 else WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.line_spacing = 1.5
    toc_styles()


def page_setup():
    for sec in doc.sections:
        sec.page_width = Cm(21)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(2.54)
        sec.bottom_margin = Cm(2.54)
        sec.left_margin = Cm(3.0)
        sec.right_margin = Cm(2.5)
        sec.header_distance = Cm(1.5)
        sec.footer_distance = Cm(1.5)


def read(name):
    return open(os.path.join(D, name), encoding="utf-8").read().splitlines()


def render_file(name, ctx):
    ctx["skip_meta"] = True
    render(read(name), ctx)


def main():
    setup_doc()
    cover()
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)

    s = open(os.path.join(D, "09-abstract-conclusion.md"), encoding="utf-8").read()
    zh = s[s.index("## 中文摘要") + len("## 中文摘要"):s.index("## Abstract")].strip()
    en = s[s.index("## Abstract") + len("## Abstract"):s.index("# 第 7 章")].strip()
    en = en.rsplit("---", 1)[0].strip()
    ch7 = s[s.index("# 第 7 章"):].strip()

    heading("摘要", 1)
    render(zh.splitlines(), {"skip_meta": False, "first_h1": False, "en": False})
    heading("Abstract", 1, page_break=True)
    render(en.splitlines(), {"skip_meta": False, "first_h1": False, "en": True})
    toc_title()
    add_toc()

    doc.add_section(WD_SECTION.NEW_PAGE)
    ctx = {"skip_meta": True, "first_h1": True, "en": False}
    for name in ("13-ch1-introduction.md", "14-ch2-literature-review.md",
                 "11-ch3-requirements.md", "12-ch4-design.md",
                 "08-ch5-implementation.md", "07-ch6-testing.md"):
        render_file(name, ctx)
    render(ch7.splitlines(), ctx)

    heading("参考文献", 1, page_break=True)
    refs = []
    for line in read("10-references.md"):
        if line.startswith("## 分布自检"):
            break
        if re.match(r"^\[\d+\]", line.strip()):
            refs.append(line.strip())
    for ref in refs:
        p = doc.add_paragraph(style="Normal")
        pf = p.paragraph_format
        pf.left_indent = Cm(0.74)
        pf.first_line_indent = Cm(-0.74)
        pf.line_spacing = 1.25
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, ref, 10.5)

    heading("致谢", 1, page_break=True)
    body("（致谢内容待补充：导师、同学与家人。）", indent=True)

    render_file("16-appendix.md", {"skip_meta": True, "first_h1": False, "en": False})

    page_setup()
    add_page_field(doc.sections[1], "lowerRoman", 1)
    add_page_field(doc.sections[2], "decimal", 1)

    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings = doc.settings.element
    compat = settings.find(qn("w:compat"))
    if compat is not None:
        compat.addprevious(update)
    else:
        settings.append(update)

    doc.save(OUT)
    print("写入", OUT)
    print("段落", len(doc.paragraphs), "表格", len(doc.tables), "图片", len(doc.inline_shapes),
          "参考文献", len(refs))


if __name__ == "__main__":
    main()
